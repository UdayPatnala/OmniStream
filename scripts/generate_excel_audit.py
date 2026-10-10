import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT_DIR = r"d:\PROJECT\AROH Open Source\Products\OmniStream"
OUTPUT_DIR = os.path.join(os.path.expanduser("~"), "Downloads")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "OmniStream_File_and_Folder_Survival_Audit.xlsx")

def classify_item(rel_path, is_dir, size_bytes):
    norm = rel_path.replace("\\", "/")
    
    # Tooling / System / External
    if norm in [".git", "node_modules", "dist"]:
        return {
            "domain": "System & Dependencies",
            "importance": "Critical" if norm == ".git" else "Redundant / Archived",
            "verdict": "NO - Must Keep" if norm == ".git" else "YES - Regenerable via npm/build",
            "desc": "Git repository history and pointers" if norm == ".git" else f"Regenerable via npm install or npm run build ({norm})",
            "risk": "Destroying version tracking" if norm == ".git" else "Zero risk - auto-recreated on build/install"
        }
    
    if norm.startswith(".agents"):
        return {
            "domain": "Agent Orchestration Logs",
            "importance": "Redundant / Archived",
            "verdict": "YES - Safe to Delete",
            "desc": "Ephemeral multi-agent coordination dispatch notes, briefings, and task progress scratch logs.",
            "risk": "Zero risk to code, runtime, tests, or build. Historical agent scratch data only."
        }
        
    if norm.startswith("v1"):
        return {
            "domain": "Legacy V1 Archive",
            "importance": "Redundant / Archived",
            "verdict": "YES - Safe to Delete",
            "desc": "Frozen legacy monolithic v1 codebase preserved prior to v1.8 OPCA refactoring.",
            "risk": "Zero risk to current application. Active modern system lives entirely in src/."
        }
        
    if norm.startswith("docs/archive"):
        return {
            "domain": "Legacy Archive",
            "importance": "Redundant / Archived",
            "verdict": "YES - Safe to Delete",
            "desc": "Legacy archive zip snapshot of OmniStream v1.5.5.",
            "risk": "Zero risk. Historical backup zip."
        }

    if norm == "records.txt":
        return {
            "domain": "Legacy Documentation",
            "importance": "Redundant / Archived",
            "verdict": "YES - Safe to Delete",
            "desc": "400-line legacy architectural symbol directory from 2024; superseded by docs/ and OMNISTREAM.md.",
            "risk": "Zero risk. Pure static text dump superseded by markdown docs."
        }

    if norm == "VERSION-CONTROL.md":
        return {
            "domain": "Version Governance",
            "importance": "Redundant / Archived",
            "verdict": "YES - Safe to Delete",
            "desc": "Duplicate secondary copy of version governance rules (superseded by VERSION_CONTROLLER.md).",
            "risk": "Zero risk. VERSION_CONTROLLER.md is the authoritative ledger."
        }

    if norm.startswith("src/app/"):
        return {
            "domain": "OPCA Shell Demo",
            "importance": "Redundant / Archived",
            "verdict": "YES - Safe to Delete",
            "desc": "Duplicate demonstration entry point created during initial OPCA refactor; active runtime uses src/App.tsx & src/main.tsx.",
            "risk": "Zero risk. App bundles from src/main.tsx and src/App.tsx."
        }

    if norm.startswith("scripts"):
        return {
            "domain": "Diagnostic & Audit Tooling",
            "importance": "Valuable",
            "verdict": "YES - Offline Tooling Only",
            "desc": "Offline CLI diagnostic scripts for manifesto compliance, test analysis, and forensic inspection.",
            "risk": "Zero risk to web app or production server; used only by developers for offline diagnostics."
        }

    if norm.startswith("docs"):
        return {
            "domain": "Product & Engineering Specs",
            "importance": "Valuable",
            "verdict": "YES - Documentation Only",
            "desc": "Authoritative system specifications, master guardian manifesto, and architectural decision records.",
            "risk": "Runtime continues to function, but architectural guidance, constitution rules, and design rationale will be lost."
        }

    if norm in ["GEMINI.md", "START_HERE.md", "OMNISTREAM.md", "ARCHITECTURE.md", "VERSION_CONTROLLER.md", "VERSION_HISTORY.md", "PROJECT_STATE.md", "HANDOFF.md", "README.md", "metadata.json", ".env.example"]:
        return {
            "domain": "Root Documentation & Constitution",
            "importance": "Valuable",
            "verdict": "YES - Reference / Constitution",
            "desc": "Constitutional rules, version ledger, onboarding ladder, and agent coordination markers.",
            "risk": "Runtime unaffected, but future developers and AI agents lose architectural continuity and decision memory."
        }

    # Configuration & Build
    if norm in ["package.json", "package-lock.json", "tsconfig.json", "vite.config.ts", "server.ts", "index.html"]:
        return {
            "domain": "Build & Runtime Foundation",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "Essential root configuration, compiler rules, dependency manifest, or server/client bootstrap.",
            "risk": "Catastrophic: Application cannot build, compile, or run."
        }

    if norm in ["vitest.config.ts", "vercel.json", "render.yaml", ".gitignore", ".github"]:
        return {
            "domain": "CI/CD & Cloud Deployment",
            "importance": "High Value",
            "verdict": "NO - CI / Cloud Config",
            "desc": "Cloud deployment rules, test runner configuration, and git version exclusion filters.",
            "risk": "Breaks automated testing, Vercel serverless routing, or Render cloud deployments."
        }

    if norm.startswith("api"):
        return {
            "domain": "Cloud Serverless API",
            "importance": "High Value",
            "verdict": "NO - Production Serverless",
            "desc": "Vercel serverless functions for YouTube search auto-complete and oEmbed video metadata proxies.",
            "risk": "Breaks YouTube search suggestions and metadata extraction when deployed to Vercel."
        }

    if norm.startswith("public"):
        is_video_or_logo = any(k in norm for k in ["mp4", "omn_logo", "cinemorph.png"])
        return {
            "domain": "Static Assets & Media",
            "importance": "Critical" if is_video_or_logo else "High Value",
            "verdict": "NO - Required Brand Asset",
            "desc": "High-fidelity brand mark, favicon, cinematic theater video intro bumper, or poster artwork plate.",
            "risk": "Broken logo/favicon, missing cinematic intro bumper, or missing fallback poster artwork."
        }

    # Tests
    if norm.startswith("src/test") or norm.startswith("src/tests"):
        return {
            "domain": "Automated Quality Assurance",
            "importance": "High Value",
            "verdict": "NO - Required for QA & Verification",
            "desc": "Automated regression tests (Tiers 1-5, OPCA boundary tests, geometry validation, 441 passing tests).",
            "risk": "Application will run, but regression protection, boundary checks, and mathematical validation are eliminated."
        }

    # Core Source Code
    if norm.startswith("src/products/cinemorph/sbh") or norm.startswith("src/lib/cinemorph/sbh"):
        return {
            "domain": "OMS-SBH Spatial Engine",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "OmniStream Spatial Behavior & Harmonization: pure geometry solver, non-overlap validation, temporal spring motion, DOM adapter.",
            "risk": "Destroys cinema screen geometry, seating non-overlap protection, and smooth spatial transitions."
        }

    if norm.startswith("src/products/cinemorph") or "cinemorph" in norm:
        return {
            "domain": "CineMorph Theater Engine",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "CineMorph theater auditorium, Web Audio 5-band DSP, adaptive ambient lighting, and ticket printing ritual.",
            "risk": "Destroys the CineMorph viewing experience and local file cinema playback."
        }

    if norm.startswith("src/products/u-tube") or "utube" in norm.lower():
        return {
            "domain": "U-Tube Discovery Engine",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "Ad-free YouTube search, discovery feed, player controls, collections, subscriptions, and history.",
            "risk": "Destroys U-Tube media discovery, watch feed, and streaming player."
        }

    if norm.startswith("src/core"):
        return {
            "domain": "OmniStream Shared Core & OMS",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "OMS decoupled intelligence pipeline, hardware capability detection, spring physics, and keyboard shortcuts.",
            "risk": "Breaks cross-engine intelligence, hardware adaptation, and responsive video stabilization."
        }

    if norm.startswith("src/shell") or norm.startswith("src/components"):
        return {
            "domain": "Shell & UI Design System",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "Global layout, header, sidebar navigation, bento cards, settings drawer, and common UI primitives.",
            "risk": "Destroys the user interface, navigation, and visual presentation."
        }

    if norm.startswith("src/state") or norm.startswith("src/services"):
        return {
            "domain": "State Management & Services",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "Zustand state stores (tickets, player, settings) and resilient localStorage persistence service.",
            "risk": "Breaks application state management, settings persistence, and ticket storage."
        }

    if norm.startswith("src/pages"):
        return {
            "domain": "Application Page Views",
            "importance": "Critical",
            "verdict": "NO - Must Keep",
            "desc": "Top-level route components for Home, Watch, Theater, Search, Channel, Collections, and Settings.",
            "risk": "Causes 404 router errors across client navigation."
        }

    # Catch-all for remaining src files
    return {
        "domain": "Application Foundation",
        "importance": "Critical",
        "verdict": "NO - Must Keep",
        "desc": "Core application module, style, type definition, or helper function.",
        "risk": "Compilation failure or runtime defect."
    }

def main():
    print(f"Scanning {ROOT_DIR}...")
    wb = openpyxl.Workbook()
    
    # ── SHEET 1: Complete Inventory ──
    ws_all = wb.active
    ws_all.title = "Complete File Inventory"
    ws_all.views.sheetView[0].showGridLines = True
    
    headers = [
        "Item Path",
        "Type",
        "Subsystem / Domain",
        "Size (KB)",
        "Importance Level",
        "Survival Verdict",
        "Functional Purpose / Description",
        "Removal Impact / Risk Assessment"
    ]
    
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws_all.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws_all.row_dimensions[1].height = 28
    
    # Colors for Importance Levels
    fills = {
        "Critical": PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid"),         # Soft Red
        "High Value": PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid"),       # Soft Amber/Yellow
        "Valuable": PatternFill(start_color="D4E6F1", end_color="D4E6F1", fill_type="solid"),         # Soft Blue
        "Redundant / Archived": PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid") # Soft Gray
    }
    
    fonts = {
        "Critical": Font(name="Calibri", size=10, bold=True, color="78281F"),
        "High Value": Font(name="Calibri", size=10, bold=True, color="7D6608"),
        "Valuable": Font(name="Calibri", size=10, bold=True, color="1B4F72"),
        "Redundant / Archived": Font(name="Calibri", size=10, bold=True, color="5D6D7E")
    }

    # Collect entries
    all_entries = []
    
    # Top level system dirs to record as single row
    system_singles = [".git", "node_modules", "dist"]
    for s in system_singles:
        p = os.path.join(ROOT_DIR, s)
        if os.path.exists(p):
            all_entries.append((s, "Directory", 0))

    for root, dirs, files in os.walk(ROOT_DIR):
        # Skip recursing into system directories
        rel_root = os.path.relpath(root, ROOT_DIR)
        parts = rel_root.replace("\\", "/").split("/")
        if any(p in system_singles for p in parts):
            continue
            
        for d in dirs:
            rel_dir = os.path.relpath(os.path.join(root, d), ROOT_DIR)
            if any(p in system_singles for p in rel_dir.replace("\\", "/").split("/")):
                continue
            all_entries.append((rel_dir, "Directory", 0))
            
        for f in files:
            full_path = os.path.join(root, f)
            rel_file = os.path.relpath(full_path, ROOT_DIR)
            try:
                size_bytes = os.path.getsize(full_path)
            except OSError:
                size_bytes = 0
            all_entries.append((rel_file, "File", size_bytes))

    # Sort entries alphabetically by path
    all_entries.sort(key=lambda x: x[0].lower())
    
    stats = {
        "Total Items": len(all_entries),
        "Total Files": sum(1 for e in all_entries if e[1] == "File"),
        "Total Directories": sum(1 for e in all_entries if e[1] == "Directory"),
        "Critical": 0,
        "High Value": 0,
        "Valuable": 0,
        "Redundant / Archived": 0,
        "Redundant Files Count": 0,
        "Redundant Size Bytes": 0
    }
    
    redundant_items = []

    current_row = 2
    for rel_path, item_type, size_bytes in all_entries:
        info = classify_item(rel_path, item_type == "Directory", size_bytes)
        imp = info["importance"]
        stats[imp] += 1
        
        size_kb = round(size_bytes / 1024, 2) if item_type == "File" else 0
        if imp == "Redundant / Archived":
            stats["Redundant Files Count"] += 1
            stats["Redundant Size Bytes"] += size_bytes
            redundant_items.append((rel_path, item_type, size_kb, info))

        row_data = [
            rel_path,
            item_type,
            info["domain"],
            size_kb if item_type == "File" else "-",
            imp,
            info["verdict"],
            info["desc"],
            info["risk"]
        ]
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_all.cell(row=current_row, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            
            # Alignments
            if col_idx in [2, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")
                
            # Importance badge styling
            if col_idx == 5:
                cell.fill = fills.get(imp, PatternFill())
                cell.font = fonts.get(imp, Font())
                
            # Verdict styling
            if col_idx == 6:
                if "Must Keep" in str(val):
                    cell.font = Font(name="Calibri", size=10, bold=True, color="78281F")
                elif "Safe to Delete" in str(val):
                    cell.font = Font(name="Calibri", size=10, bold=True, color="1E8449")
                elif "CI" in str(val):
                    cell.font = Font(name="Calibri", size=10, bold=True, color="7D6608")

        ws_all.row_dimensions[current_row].height = 20
        current_row += 1

    # ── SHEET 2: Pruning & Cleanup Candidates ──
    ws_prune = wb.create_sheet(title="Safe to Prune")
    ws_prune.views.sheetView[0].showGridLines = True
    
    prune_headers = [
        "Item Path",
        "Type",
        "Subsystem / Domain",
        "Size (KB)",
        "Why Safe to Delete?",
        "Impact on Project Survival"
    ]
    
    prune_fill = PatternFill(start_color="1E8449", end_color="1E8449", fill_type="solid")
    for col_idx, h in enumerate(prune_headers, 1):
        cell = ws_prune.cell(row=1, column=col_idx, value=h)
        cell.fill = prune_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_prune.row_dimensions[1].height = 28
    
    p_row = 2
    for r_path, r_type, r_kb, r_info in redundant_items:
        r_vals = [
            r_path,
            r_type,
            r_info["domain"],
            r_kb if r_type == "File" else "-",
            r_info["desc"],
            r_info["risk"]
        ]
        for col_idx, val in enumerate(r_vals, 1):
            cell = ws_prune.cell(row=p_row, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if col_idx in [2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")
        ws_prune.row_dimensions[p_row].height = 20
        p_row += 1

    # ── SHEET 3: Executive Summary & Dashboard ──
    ws_sum = wb.create_sheet(title="Executive Summary", index=0)
    ws_sum.views.sheetView[0].showGridLines = True
    
    # Title
    ws_sum.merge_cells("A1:E1")
    title_cell = ws_sum["A1"]
    title_cell.value = "OmniStream / CineMorph Project Directory Forensic Audit & Survival Assessment"
    title_cell.fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 36
    
    summary_data = [
        ("Metric / Assessment", "Count", "Percentage", "Survival Verdict", "Explanation"),
        ("Total Scanned Items", stats["Total Items"], "100%", "Full Inventory", "All files and folders across root, src, public, docs, scripts, etc."),
        ("Total Files", stats["Total Files"], f"{round(stats['Total Files']/stats['Total Items']*100, 1)}%", "-", "Actual files on disk (excluding node_modules internal files)."),
        ("Total Directories", stats["Total Directories"], f"{round(stats['Total Directories']/stats['Total Items']*100, 1)}%", "-", "Directory structures defining architecture."),
        ("", "", "", "", ""),
        ("🔴 Critical (Must Keep)", stats["Critical"], f"{round(stats['Critical']/stats['Total Items']*100, 1)}%", "CANNOT SURVIVE WITHOUT", "Core application code, main entry, build config, stores, and primary pages."),
        ("🟡 High Value (Verification / Cloud)", stats["High Value"], f"{round(stats['High Value']/stats['Total Items']*100, 1)}%", "REQUIRED FOR QA & DEPLOY", "Vitest automated suites (76 suites, 441 tests), OPCA boundaries, and Vercel serverless functions."),
        ("🔵 Valuable (Documentation & Tooling)", stats["Valuable"], f"{round(stats['Valuable']/stats['Total Items']*100, 1)}%", "SURVIVES WITHOUT (KNOWLEDGE LOSS)", "Architecture specs, constitution, version ledger, and offline diagnostic scripts."),
        ("⚪ Redundant / Archived (Safe to Delete)", stats["Redundant / Archived"], f"{round(stats['Redundant / Archived']/stats['Total Items']*100, 1)}%", "SURVIVES EASILY (CLEANUP TARGET)", ".agents/ coordination scratch logs, v1/ legacy backup, archive zip, and duplicate entry points."),
        ("", "", "", "", ""),
        ("Cleanable Space Estimate", f"{round(stats['Redundant Size Bytes'] / (1024*1024), 2)} MB", "-", "Safe Pruning Potential", f"{stats['Redundant Files Count']} files and redundant folders can be safely removed without affecting code or tests.")
    ]
    
    for r_idx, row in enumerate(summary_data, 3):
        ws_sum.row_dimensions[r_idx].height = 24
        for c_idx, val in enumerate(row, 1):
            cell = ws_sum.cell(row=r_idx, column=c_idx, value=val)
            if r_idx == 3:
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = Font(name="Calibri", size=10, bold=(c_idx in [1, 4]))
                cell.border = thin_border
                if c_idx in [2, 3]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(vertical="center")
                    
                # Highlight verdict
                if "CANNOT SURVIVE" in str(val):
                    cell.fill = fills["Critical"]
                    cell.font = fonts["Critical"]
                elif "REQUIRED FOR QA" in str(val):
                    cell.fill = fills["High Value"]
                    cell.font = fonts["High Value"]
                elif "SURVIVES EASILY" in str(val):
                    cell.fill = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="1E8449")

    # Auto-adjust column widths for all sheets
    for ws in [ws_sum, ws_all, ws_prune]:
        ws.freeze_panes = "A2" if ws != ws_sum else "A4"
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len and len(val) < 80:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(f"Excel audit successfully saved to: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
